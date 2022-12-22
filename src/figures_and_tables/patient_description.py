"""Produce table of sex and age distribution along with diagnosis"""
from pathlib import Path

import numpy as np
import pandas as pd
from psycopmlutils.sql.loader import sql_load

import docx


def table_to_docx(df, path):
    doc = docx.Document()
    t = doc.add_table(df.shape[0] + 1, df.shape[1])

    # add the header rows.
    for j in range(df.shape[-1]):
        t.cell(0, j).text = df.columns[j]

    # add the rest of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i + 1, j).text = str(df.values[i, j])

    # save the doc
    doc.save(path)


if __name__ == "__main__":

    TABLE_NAME = "[FOR_kohorte_indhold_pt_journal_inkl_2021_feb2022]"
    df = sql_load(
        query=f"SELECT * FROM [FCT].{TABLE_NAME}",
        format_timestamp_cols_to_datetime=True,
    )

    df["koennavn"] = np.where(df["koennavn"] == "Mand", "Male", "Female")
    df["age_int"] = df["min_alder"].apply(lambda x: int(np.floor(x)))
    df["age_bin"] = pd.cut(df["age_int"], bins=[0, 18, 30, 50, 70, 110])
    # Some contacts don't have an ending time (which we need to find the most recent admission)
    # For those, setting ending ending time to start time
    df["datotid_slut"] = df["datotid_slut"].fillna(df["datotid_start"])
    # Find most recent ending time
    df["max_dato_slut"] = df.groupby("dw_ek_borger")["datotid_slut"].transform(max)
    ## Find most recent a-diagnosis
    # Get most recent visit
    most_recent = df[df["max_dato_slut"] == df["datotid_slut"]]
    # Some have more than one contact registered at the maximum end time
    # if so, choose the contact that starts the latest
    most_recent = most_recent[
        most_recent.groupby("dw_ek_borger")["datotid_start"].transform(max)
        == most_recent["datotid_start"]
    ]
    # still ~50 duplicates. Keeping the last occurence
    most_recent = most_recent.drop_duplicates(
        subset=["dw_ek_borger", "datotid_slut", "datotid_start"], keep="last"
    )
    most_recent = most_recent[["dw_ek_borger", "diagnosis"]]
    most_recent.loc[
        most_recent["diagnosis"] == "F99-F99 Unspecified mental disorder", "diagnosis"
    ] = "Others"
    most_recent = most_recent.rename({"diagnosis": "last_diagnosis"}, axis=1)
    # Adding final a-diagnosis to main df
    df = df.merge(most_recent, on="dw_ek_borger", validate="many_to_one")

    sex_age_df = (
        df.groupby(["last_diagnosis", "age_group", "koennavn"])["dw_ek_borger"]
        .nunique()
        .reset_index()
        .pivot(index="last_diagnosis", columns=["age_group", "koennavn"])
    )
    sex_age_df_bin = (
        df.groupby(["last_diagnosis", "age_bin", "koennavn"])["dw_ek_borger"]
        .nunique()
        .reset_index()
        .pivot(index="last_diagnosis", columns=["age_bin", "koennavn"])
    )

    # Add sum row and column
    # sex_age_df["Total"] = sex_age_df.sum(axis=1)
    # sex_age_df.loc["Total",:] = sex_age_df.sum(axis=0)

    save_dir = Path(__file__).parent.parent.parent / "tables"
    if not save_dir.exists():
        save_dir.mkdir()
    table_to_docx(sex_age_df, save_dir / "demographics.docx")
    table_to_docx(sex_age_df.reset_index(), save_dir / "demographics.docx")
    table_to_docx(sex_age_df_bin, save_dir / "demographics_bins.docx")
    table_to_docx(sex_age_df_bin.reset_index(), save_dir / "demographics_bins.docx")
